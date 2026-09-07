"""Geometry QA. LibreOffice is unavailable here, so bounds, overlap and text
overflow are checked numerically instead of by eye."""
import math, sys
from pptx import Presentation
from pptx.util import Emu

EMU = 914400.0
prs = Presentation("eligibility-compiler.pptx")
SW, SH = prs.slide_width / EMU, prs.slide_height / EMU
print(f"canvas {SW:.2f} x {SH:.2f} in, {len(prs.slides.__iter__.__self__._sldIdLst)} slides\n")

# average glyph width as a fraction of point size, per family
WIDTH = {"Cambria": 0.50, "Calibri": 0.47, None: 0.50}
issues = []

def est_lines(text, width_in, pt, face):
    if not text.strip():
        return 0
    per_char = pt * WIDTH.get(face, 0.5) / 72.0
    if per_char <= 0:
        return 1
    cap = max(1, int(width_in / per_char))
    lines = 0
    for para in text.split("\n"):
        lines += max(1, math.ceil(len(para) / cap))
    return lines

for i, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        if sh.left is None:
            continue
        L, T = sh.left / EMU, sh.top / EMU
        Wd, Ht = (sh.width or 0) / EMU, (sh.height or 0) / EMU
        R, Bm = L + Wd, T + Ht

        if L < -0.02 or T < -0.02 or R > SW + 0.02 or Bm > SH + 0.02:
            issues.append(f"s{i}: OFF-SLIDE {sh.shape_type} "
                          f"[{L:.2f},{T:.2f},{R:.2f},{Bm:.2f}]")
        elif L < 0.45 or T < 0.12 or R > SW - 0.45 or Bm > SH - 0.12:
            txt = (sh.text_frame.text[:26] if sh.has_text_frame else "")
            issues.append(f"s{i}: tight margin [{L:.2f},{T:.2f},{R:.2f},"
                          f"{Bm:.2f}] {txt!r}")

        if sh.has_text_frame and sh.text_frame.text.strip():
            t = sh.text_frame.text
            pt, face = 12, None
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.font.size:
                        pt = max(pt, r.font.size.pt)
                    if r.font.name:
                        face = r.font.name
            n = est_lines(t, Wd - 0.1, pt, face)
            need = n * pt * 1.30 / 72.0
            if need > Ht + 0.06:
                issues.append(f"s{i}: OVERFLOW ~{need:.2f}in needed vs "
                              f"{Ht:.2f}in box, {pt:.0f}pt — {t[:44]!r}")
            boxes.append((L, T, R, Bm, t[:26]))

    for a in range(len(boxes)):
        for b in range(a + 1, len(boxes)):
            ax, ay, ar, ab, at = boxes[a]
            bx, by, br, bb, bt = boxes[b]
            ox = min(ar, br) - max(ax, bx)
            oy = min(ab, bb) - max(ay, by)
            if ox > 0.10 and oy > 0.10:
                issues.append(f"s{i}: TEXT OVERLAP {ox:.2f}x{oy:.2f}in "
                              f"{at!r} / {bt!r}")

if issues:
    print(f"{len(issues)} issue(s):")
    for x in issues:
        print("  " + x)
else:
    print("no geometry issues found")
